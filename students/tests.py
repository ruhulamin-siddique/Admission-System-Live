from io import BytesIO

import pandas as pd
from django.contrib.auth.models import User
from django.test import TestCase, override_settings
from django.urls import reverse

from core.models import Role, RolePermission
from master_data.models import Cluster, Program
from students.models import Student


@override_settings(ALLOWED_HOSTS=['testserver', 'localhost'])
class StudentDirectoryTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        from students.utils import get_canonical_program_name
        if hasattr(get_canonical_program_name, '_cache'):
            delattr(get_canonical_program_name, '_cache')
            
        engineering = Cluster.objects.create(name='Engineering & Technology', code='05')
        business = Cluster.objects.create(name='Business & Management', code='09')

        from master_data.models import Hall
        Hall.objects.get_or_create(short_name='AUAH', code='01')
        Hall.objects.get_or_create(short_name='BTBH', code='02')
        Hall.objects.get_or_create(short_name='ZH', code='03')

        cls.cse_program = Program.objects.create(
            name='Computer Science and Engineering',
            short_name='CSE',
            ugc_code='01',
            cluster=engineering,
            level_code='1',
        )
        cls.eee_program = Program.objects.create(
            name='Electrical and Electronic Engineering',
            short_name='EEE',
            ugc_code='02',
            cluster=engineering,
            level_code='1',
        )
        Program.objects.create(
            name='Master of Business Administration',
            short_name='MBA',
            ugc_code='03',
            cluster=business,
            level_code='3',
        )

        for index in range(1, 16):
            Student.objects.create(
                student_id=f'CSE{index:03d}',
                student_name=f'CSE Student {index:02d}',
                program=cls.cse_program.name,
                admission_year=2025 if index <= 10 else 2024,
                cluster=engineering.name,
                batch='25th' if index <= 10 else '24th',
                semester_name='Spring',
                program_type='Bachelor',
                admission_status='Active',
                gender='Male' if index % 2 else 'Female',
                father_name=f'Father CSE {index:02d}',
                mother_name=f'Mother CSE {index:02d}',
                student_mobile=f'01700000{index:03d}',
            )

        cls.inactive_student = Student.objects.create(
            student_id='CSE900',
            student_name='Inactive CSE Student',
            program=cls.cse_program.name,
            admission_year=2024,
            cluster=engineering.name,
            batch='24th',
            semester_name='Fall',
            program_type='Bachelor',
            admission_status='Inactive',
            gender='Female',
            father_name='Inactive Father',
            mother_name='Inactive Mother',
            father_mobile='01711000001',
            mother_mobile='01711000002',
            student_mobile='01810000000',
            student_email='inactive@example.com',
            emergency_contact='01711999999',
            blood_group='B+',
            religion='Islam',
            present_address='Dormitory Road, Section 1, Dhaka Cantonment',
            permanent_address='Village Home, Cumilla Sadar, Cumilla',
            photo_path='photos/inactive.jpg',
        )
        cls.cancelled_student = Student.objects.create(
            student_id='CSE901',
            student_name='Cancelled CSE Student',
            program=cls.cse_program.name,
            admission_year=2025,
            cluster=engineering.name,
            batch='25th',
            semester_name='Spring',
            program_type='Bachelor',
            admission_status='Cancelled',
            gender='Male',
            father_name='Cancelled Father',
            mother_name='Cancelled Mother',
            student_mobile='01810000001',
            photo_path='https://example.com/cancelled.jpg',
        )
        cls.legacy_scope_student = Student.objects.create(
            student_id='CSE902',
            student_name='Legacy Short Program Student',
            program='CSE',
            admission_year=2025,
            cluster=engineering.name,
            batch='25th',
            semester_name='Spring',
            program_type='Bachelor',
            admission_status='Active',
            gender='Male',
            father_name='Legacy Father',
            mother_name='Legacy Mother',
            student_mobile='01810000002',
        )

        for index in range(1, 5):
            Student.objects.create(
                student_id=f'EEE{index:03d}',
                student_name=f'EEE Student {index:02d}',
                program=cls.eee_program.name,
                admission_year=2025,
                cluster=engineering.name,
                batch='25th',
                semester_name='Spring',
                program_type='Bachelor',
                admission_status='Active',
                gender='Female',
                father_name=f'Father EEE {index:02d}',
                mother_name=f'Mother EEE {index:02d}',
                student_mobile=f'01900000{index:03d}',
            )

        cls.latest_eee_student = Student.objects.create(
            student_id='EEE900',
            student_name='EEE Latest Batch Student',
            program=cls.eee_program.name,
            admission_year=2026,
            cluster=engineering.name,
            batch='26th',
            semester_name='Spring',
            program_type='Bachelor',
            admission_status='Active',
            gender='Female',
            father_name='Latest EEE Father',
            mother_name='Latest EEE Mother',
            student_mobile='01999999999',
        )

        cls.mba_student = Student.objects.create(
            student_id='MBA001',
            student_name='MBA Student',
            program='MBA',
            admission_year=2025,
            cluster=business.name,
            batch='MBA-1',
            semester_name='Spring',
            program_type='Masters',
            admission_status='Active',
            gender='Female',
            father_name='MBA Father',
            mother_name='MBA Mother',
            student_mobile='01610000000',
        )

        cls.superuser = User.objects.create_superuser(
            username='admin',
            email='admin@example.com',
            password='password123',
        )

        scoped_role = Role.objects.create(name='Scoped Staff')
        RolePermission.objects.create(role=scoped_role, module='students', task='view_directory')
        RolePermission.objects.create(role=scoped_role, module='students', task='export_excel')

        cls.scoped_user = User.objects.create_user(
            username='scoped',
            email='scoped@example.com',
            password='password123',
        )
        cls.scoped_user.profile.role = scoped_role
        cls.scoped_user.profile.department_scope = 'CSE'
        cls.scoped_user.profile.save()

    def setUp(self):
        self.client.force_login(self.superuser)

    def read_excel(self, response):
        return pd.read_excel(BytesIO(response.content))

    def test_directory_excludes_cancelled_by_default(self):
        response = self.client.get(reverse('student_list'))

        self.assertEqual(response.status_code, 200)
        self.assertNotIn(self.cancelled_student, response.context['page_obj'].paginator.object_list)
        self.assertContains(response, 'Showing all students except')

    def test_directory_status_filter_can_show_cancelled(self):
        response = self.client.get(reverse('student_list'), {'status': 'Cancelled'})

        students = list(response.context['page_obj'].object_list)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(students, [self.cancelled_student])

    def test_directory_renders_combined_parent_column(self):
        response = self.client.get(
            reverse('student_list'),
            {'search': self.inactive_student.student_id},
            HTTP_HX_REQUEST='true',
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, '<th>Parents</th>', html=True)
        self.assertContains(response, f'Father: {self.inactive_student.father_name}')
        self.assertContains(response, f'Mother: {self.inactive_student.mother_name}')

    def test_directory_page_has_clear_filters_button(self):
        response = self.client.get(reverse('student_list'))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'btn-clear-filters')
        self.assertContains(response, 'name="sort"')
        self.assertContains(response, '?per_page=25')
        self.assertContains(response, 'sort=batch_dept_serial')

    def test_directory_filters_work_for_visible_fields(self):
        cases = (
            ({'search': 'Inactive CSE Student'}, lambda rows: rows == [self.inactive_student]),
            ({'year': '2024'}, lambda rows: all(student.admission_year == 2024 for student in rows)),
            ({'dept': 'Business & Management'}, lambda rows: rows == [self.mba_student]),
            ({'program': 'CSE'}, lambda rows: all(student.program == 'CSE' for student in rows)),
            ({'batch': 'MBA-1'}, lambda rows: rows == [self.mba_student]),
            ({'type': 'Masters'}, lambda rows: rows == [self.mba_student]),
            ({'gender': 'Female'}, lambda rows: rows and all(student.gender == 'Female' for student in rows)),
        )

        for params, assertion in cases:
            with self.subTest(params=params):
                response = self.client.get(
                    reverse('student_list'),
                    params,
                    HTTP_HX_REQUEST='true',
                )
                rows = list(response.context['page_obj'].paginator.object_list)
                self.assertEqual(response.status_code, 200)
                self.assertTrue(assertion(rows))

    def test_default_directory_sort_is_batch_then_department_then_serial(self):
        response = self.client.get(reverse('student_list'))

        ordered_ids = [student.student_id for student in response.context['page_obj'].paginator.object_list]

        self.assertEqual(response.status_code, 200)
        self.assertEqual(ordered_ids[0], self.latest_eee_student.student_id)
        self.assertLess(ordered_ids.index('CSE001'), ordered_ids.index('EEE001'))

    def test_department_sort_prioritizes_department_before_batch(self):
        response = self.client.get(reverse('student_list'), {'sort': 'dept_batch_serial'})

        ordered_ids = [student.student_id for student in response.context['page_obj'].paginator.object_list]

        self.assertEqual(response.status_code, 200)
        self.assertEqual(ordered_ids[:3], ['CSE001', 'CSE002', 'CSE003'])
        self.assertLess(ordered_ids.index('CSE010'), ordered_ids.index('CSE011'))
        self.assertLess(ordered_ids.index('CSE900'), ordered_ids.index('EEE900'))

    def test_htmx_directory_results_include_synced_controls(self):
        response = self.client.get(
            reverse('student_list'),
            {'program': self.cse_program.name, 'per_page': 10},
            HTTP_HX_REQUEST='true',
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'hx-target="#directory-results"')
        self.assertContains(response, 'hx-include="#directory-filter-form"')
        self.assertContains(response, 'id="total-count-display"')
        self.assertContains(response, reverse('export_students_all'))

    def test_directory_pagination_and_page_size_preserve_filters(self):
        response = self.client.get(
            reverse('student_list'),
            {'program': 'CSE', 'page': 2, 'per_page': 10},
            HTTP_HX_REQUEST='true',
        )

        page_obj = response.context['page_obj']
        self.assertEqual(response.status_code, 200)
        self.assertEqual(page_obj.number, 2)
        self.assertEqual(page_obj.paginator.per_page, 10)
        self.assertTrue(all(student.program == 'CSE' for student in page_obj.object_list))
        self.assertContains(response, 'directory-per-page-control')
        self.assertContains(response, 'hx-get="/students/?page=1"')

    def test_htmx_pagination_footer_hides_filtered_count_sentence(self):
        response = self.client.get(
            reverse('student_list'),
            {'program': self.cse_program.name},
            HTTP_HX_REQUEST='true',
        )

        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, 'Showing <strong>')

    def test_api_preview_id_requires_add_student_permission(self):
        self.client.force_login(self.scoped_user)

        response = self.client.get(reverse('api_preview_id'))

        self.assertRedirects(response, reverse('user_profile'))

    def test_bulk_update_rejects_unknown_field(self):
        response = self.client.post(reverse('api_bulk_update_execute'), {
            'student_ids': self.inactive_student.student_id,
            'field_name': 'student_name',
            'new_value': 'Changed By Bulk Update',
        })

        self.assertEqual(response.status_code, 400)
        self.assertContains(response, 'cannot be updated in bulk', status_code=400)
        self.inactive_student.refresh_from_db()
        self.assertEqual(self.inactive_student.student_name, 'Inactive CSE Student')

    def test_scoped_user_sees_and_exports_mapped_programs(self):
        self.client.force_login(self.scoped_user)

        response = self.client.get(reverse('student_list'))
        visible_programs = {student.program for student in response.context['page_obj'].paginator.object_list}

        self.assertEqual(response.status_code, 200)
        self.assertEqual(visible_programs, {'CSE'})

        export_response = self.client.get(reverse('export_students'))
        export_df = self.read_excel(export_response)
        self.assertEqual(set(export_df['program']), {'CSE'})
        self.assertEqual(len(export_df), 17)

    def test_standard_export_matches_filtered_directory_dataset(self):
        params = {'program': 'CSE', 'status': 'Inactive'}
        list_response = self.client.get(reverse('student_list'), params)
        export_response = self.client.get(reverse('export_students'), params)

        self.assertEqual(list_response.context['total_count'], 1)
        export_df = self.read_excel(export_response)
        self.assertEqual(len(export_df), 1)
        self.assertEqual(export_df.iloc[0]['student_id'], self.inactive_student.student_id)

    def test_export_respects_selected_sort_order(self):
        response = self.client.get(reverse('export_students'), {'sort': 'batch_dept_serial'})

        export_df = self.read_excel(response)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(export_df.iloc[0]['student_id'], self.latest_eee_student.student_id)

    def test_all_info_export_includes_all_student_fields(self):
        response = self.client.get(reverse('export_students_all'), {'search': self.inactive_student.student_id})

        export_df = self.read_excel(response)
        expected_columns = [field.name for field in Student._meta.concrete_fields]

        self.assertEqual(response.status_code, 200)
        self.assertEqual(list(export_df.columns), expected_columns)
        self.assertEqual(len(export_df), 1)
        self.assertEqual(export_df.iloc[0]['student_id'], self.inactive_student.student_id)

    def test_student_photo_url_handles_relative_and_absolute_paths(self):
        self.assertEqual(self.inactive_student.photo_url, '/media/photos/inactive.jpg')
        self.assertEqual(self.cancelled_student.photo_url, 'https://example.com/cancelled.jpg')

    def test_quick_info_card_shows_only_extra_details_not_visible_in_rows(self):
        response = self.client.get(reverse('student_short_info', args=[self.inactive_student.student_id]))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Quick Student Details')
        self.assertContains(response, self.inactive_student.student_email)
        self.assertContains(response, self.inactive_student.father_mobile)
        self.assertContains(response, self.inactive_student.mother_mobile)
        self.assertContains(response, self.inactive_student.emergency_contact)
        self.assertContains(response, self.inactive_student.present_address)
        self.assertContains(response, self.inactive_student.permanent_address)
        self.assertContains(response, self.inactive_student.blood_group)
        self.assertContains(response, self.inactive_student.religion)
        self.assertNotContains(response, self.inactive_student.father_name)
        self.assertNotContains(response, self.cse_program.name)
        self.assertNotContains(response, 'Open Profile')

    def test_institutional_report_access_and_rendering(self):
        # Update some students to have ssc_school and hsc_college
        student1 = Student.objects.get(student_id='CSE001')
        student1.ssc_school = 'Lions School'
        student1.hsc_college = 'Lions College'
        student1.save()

        student2 = Student.objects.get(student_id='CSE002')
        student2.ssc_school = 'Lions School'
        student2.hsc_college = 'Blue College'
        student2.save()

        # Check access for superuser
        response = self.client.get(reverse('institutional_report'), {'year': '2025'})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Lions School')
        self.assertContains(response, 'Lions College')
        self.assertContains(response, 'Blue College')

        # Check access denied for scoped_user (no view_analytics under reports)
        self.client.force_login(self.scoped_user)
        response = self.client.get(reverse('institutional_report'))
        self.assertRedirects(response, reverse('user_profile'))

    def test_api_institutional_students_modal_data(self):
        student1 = Student.objects.get(student_id='CSE001')
        student1.ssc_school = 'Lions School'
        student1.save()

        # Query API for students from Lions School
        response = self.client.get(reverse('api_institutional_students'), {'school': 'Lions School', 'year': '2025'})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'CSE001')
        self.assertContains(response, 'Lions School')
        # Print URL should be formatted in response
        self.assertContains(response, reverse('print_institutional_students'))

    def test_print_institutional_students_view(self):
        student1 = Student.objects.get(student_id='CSE001')
        student1.ssc_school = 'Lions School'
        student1.save()

        response = self.client.get(reverse('print_institutional_students'), {'school': 'Lions School'})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'CSE001')
        self.assertContains(response, 'Lions School')
        self.assertContains(response, 'window.print()')

    def test_dashboard_reference_nodes(self):
        # Create ReferenceNode objects
        from students.models import ReferenceNode
        ref_fb = ReferenceNode.objects.create(name_en='Facebook Ad')
        ref_alumni = ReferenceNode.objects.create(name_en='Alumni Network')

        # Update some students to have references
        student1 = Student.objects.get(student_id='CSE001')
        student1.reference = ref_fb
        student1.batch = '26th'
        student1.save()

        student2 = Student.objects.get(student_id='CSE002')
        student2.reference = ref_alumni
        student2.batch = '26th'
        student2.save()

        student3 = Student.objects.get(student_id='CSE003')
        student3.reference = ref_fb
        student3.batch = '26th'
        student3.save()

        # Query the dashboard
        response = self.client.get(reverse('dashboard'))
        self.assertEqual(response.status_code, 200)

        # Assert top_references is in context
        stats = response.context['stats']
        self.assertIn('top_references', stats)
        
        # Verify Facebook Ad is ranked #1 (since it has 2 count, Alumni has 1)
        top_refs = stats['top_references']
        self.assertEqual(top_refs[0]['reference'], 'Facebook Ad')
        self.assertEqual(top_refs[0]['count'], 2)
        self.assertEqual(top_refs[1]['reference'], 'Alumni Network')
        self.assertEqual(top_refs[1]['count'], 1)

        # Assert page rendering has the reference text
        self.assertContains(response, 'Admission Reference')
        self.assertContains(response, 'Facebook Ad')
        self.assertContains(response, 'Alumni Network')

    def test_student_field_audit_and_revert(self):
        from students.models import StudentFieldHistory
        
        # 1. Edit a student's profile and confirm changes are logged
        student = Student.objects.get(student_id='CSE001')
        self.assertEqual(student.student_name, 'CSE Student 01')
        
        student.student_name = 'Updated Name'
        student.changed_by_user = self.superuser
        student.save()
        
        # Verify history entry was created
        history = StudentFieldHistory.objects.filter(student=student, field_name='student_name')
        self.assertEqual(history.count(), 1)
        entry = history.first()
        self.assertEqual(entry.old_value, 'CSE Student 01')
        self.assertEqual(entry.new_value, 'Updated Name')
        self.assertEqual(entry.changed_by, self.superuser)
        self.assertFalse(entry.reverted)
        
        # 2. Test permission checks on the revert API
        revert_url = reverse('revert_field_change', args=[entry.id])
        
        # Non-superuser should be denied (returns 403)
        self.client.force_login(self.scoped_user)
        resp_scoped = self.client.post(revert_url)
        self.assertEqual(resp_scoped.status_code, 403)
        
        # Superuser should succeed (returns 200 JSON with success: True)
        self.client.force_login(self.superuser)
        resp_super = self.client.post(revert_url)
        self.assertEqual(resp_super.status_code, 200)
        self.assertTrue(resp_super.json()['success'])
        
        # Verify database reverted to old value
        student.refresh_from_db()
        self.assertEqual(student.student_name, 'CSE Student 01')
        
        # Verify history entry state updated
        entry.refresh_from_db()
        self.assertTrue(entry.reverted)
        self.assertEqual(entry.reverted_by, self.superuser)
        self.assertIsNotNone(entry.reverted_at)

    def test_dashboard_batch_department_matrix(self):
        # Request the dashboard page
        response = self.client.get(reverse('dashboard'))
        self.assertEqual(response.status_code, 200)

        # Assert context variables are present
        self.assertIn('matrix_cols', response.context)
        self.assertIn('matrix_rows', response.context)
        self.assertIn('column_totals', response.context)
        self.assertIn('grand_total', response.context)

        cols = response.context['matrix_cols']
        rows = response.context['matrix_rows']

        # Assert column headers (programs) are sorted alphabetically
        self.assertEqual(cols, sorted(cols))

        # Assert rows (batches) are sorted descending by batch_number
        prev_num = float('inf')
        for r in rows:
            batch_num = r['batch_number'] or 0
            self.assertTrue(batch_num <= prev_num, f"Batch {r['batch_name']} with number {batch_num} is out of order (previous was {prev_num})")
            prev_num = batch_num

        # Assert grand total sums up all students in database
        self.assertEqual(response.context['grand_total'], Student.objects.count())

    def test_ugc_id_generation_program_lookup_exact(self):
        from students.utils import generate_next_ugc_id, generate_ugc_prefix
        from master_data.models import Cluster, Program

        # 1. Create a dummy Arts cluster and English program
        arts_cluster = Cluster.objects.create(name='Arts', code='02')
        english_program = Program.objects.create(
            name='English',
            short_name='ENG',
            ugc_code='09',
            cluster=arts_cluster,
            level_code='1'
        )

        # 2. Create a dummy Civil Engineering program that could clash if contains matches
        civil_program = Program.objects.create(
            name='Civil Engineering',
            short_name='CE',
            ugc_code='04',
            cluster=self.cse_program.cluster,  # Engineering cluster code '05'
            level_code='1'
        )

        # Test generate_next_ugc_id for program_name='ENG'
        # Arts code = 02, English ugc_code = 09, Level code = 1. Part CcssP should be 02091
        next_id = generate_next_ugc_id(
            admission_year=2026,
            semester_name='Spring',
            hall_name='Non-Residential',
            program_name='ENG',
            cluster_name='Arts'
        )
        # UGC ID prefix: UUU(080) YY(26) S(1) HH(00) CcssP(02091)
        expected_prefix = "0802610002091"
        self.assertTrue(next_id.startswith(expected_prefix), f"Expected ID to start with {expected_prefix}, got {next_id}")

        # Test generate_ugc_prefix for program_name='ENG'
        prefix = generate_ugc_prefix(
            admission_year=2026,
            semester_name='Spring',
            hall_name='Non-Residential',
            program_name='ENG',
            cluster_name='Arts'
        )
        self.assertEqual(prefix, expected_prefix)

    def test_excel_import_logs_activity(self):
        from core.models import ActivityLog
        # Create Excel file in-memory
        data = {
            'student_id': ['0802510001011001'],
            'student_name': ['IMPORT STUDENT ONE'],
            'program': ['CSE'],
            'admission_status': ['Active'],
            'gender': ['Male']
        }
        df = pd.DataFrame(data)
        excel_file = BytesIO()
        df.to_excel(excel_file, index=False)
        excel_file.seek(0)
        excel_file.name = 'test_import.xlsx'

        # Perform POST request to import_students
        import_url = reverse('import_students')
        response = self.client.post(import_url, {'excel_file': excel_file})
        self.assertEqual(response.status_code, 200)

        # Check that Student was created
        student_exists = Student.objects.filter(student_id='0802510001011001').exists()
        self.assertTrue(student_exists)

        # Check ActivityLog entry
        log_entry = ActivityLog.objects.filter(module='students', object_id='0802510001011001', action_type='CREATE').first()
        self.assertIsNotNone(log_entry)
        self.assertEqual(log_entry.user, self.superuser)
        self.assertIn('Imported student record', log_entry.description)

        # Now test the timeline rendering of the imported student profile view!
        profile_url = reverse('student_profile', args=['0802510001011001'])
        profile_response = self.client.get(profile_url)
        self.assertEqual(profile_response.status_code, 200)
        timeline = profile_response.context['timeline']
        
        # Verify timeline has correct import details
        import_event = next((e for e in timeline if 'Imported' in e['title'] or 'Record Imported' in e['title']), None)
        self.assertIsNotNone(import_event)
        self.assertEqual(import_event['user'], self.superuser.username)
        self.assertIn('imported into the system via Excel spreadsheet', import_event['description'])

    def test_api_matrix_students_drilldown(self):
        # We query the api_matrix_students endpoint for CSE and batch 25th
        matrix_url = reverse('api_matrix_students')
        
        # Test specific batch and program filter
        response = self.client.get(matrix_url, {'batch': '25th', 'program': 'CSE'})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Showing')
        self.assertContains(response, 'CSE Student')
        
        # Test total column drill-down (all programs in batch)
        response_batch_total = self.client.get(matrix_url, {'batch': '25th', 'program': 'Total'})
        self.assertEqual(response_batch_total.status_code, 200)
        self.assertContains(response_batch_total, 'CSE Student')
        self.assertContains(response_batch_total, 'EEE Student')
        
        # Test total row drill-down (all batches in program)
        response_program_total = self.client.get(matrix_url, {'batch': 'Total', 'program': 'CSE'})
        self.assertEqual(response_program_total.status_code, 200)
        self.assertContains(response_program_total, 'CSE Student')
        
        # Test grand total drill-down
        response_grand_total = self.client.get(matrix_url, {'batch': 'Total', 'program': 'Total'})
        self.assertEqual(response_grand_total.status_code, 200)
        self.assertContains(response_grand_total, 'CSE Student')
        self.assertContains(response_grand_total, 'EEE Student')
        self.assertContains(response_grand_total, 'MBA Student')

    def test_legacy_student_form_validation(self):
        from students.forms import StudentForm
        form_data = {
            'student_name': 'Legacy test student',
            'is_legacy_student': 'on',
            'old_student_id': '123456789',
            'program': self.cse_program.name,
            'admission_year': 2024,
            'semester_name': 'Spring',
            'program_type': 'Bachelor',
            'batch': '24th',
            'current_batch': '24th',
            'current_semester': 'Level 1 Term I',
            'student_mobile': '01700000000',
            'father_mobile': '01700000001',
            'mother_mobile': '01700000002',
            'ssc_school': 'SSC School',
            'ssc_year': '2022',
            'ssc_board': 'Dhaka',
            'ssc_roll': '123456',
            'ssc_gpa': 5.00,
            'hsc_college': 'HSC College',
            'hsc_year': '2024',
            'hsc_board': 'Dhaka',
            'hsc_roll': '123456',
            'hsc_gpa': 5.00,
            'admission_payment': 10000,
            'admission_status': 'Active',
        }
        form = StudentForm(data=form_data)
        self.assertTrue(form.is_valid(), form.errors.as_json())
        self.assertIsNotNone(form.cleaned_data['student_id'])
        self.assertTrue(form.cleaned_data['student_id'].startswith('080'))

    def test_api_student_ugc_id_preview(self):
        # Create a legacy student with missing academic parameters
        legacy_student_missing = Student.objects.create(
            student_id='999999',
            student_name='Legacy Missing Fields',
            is_legacy_student=True,
            admission_year=None,
            semester_name=None,
            hall_attached=None,
            program=None,
            cluster=None
        )
        
        url = reverse('api_student_ugc_id_preview', kwargs={'student_id': legacy_student_missing.student_id})
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertFalse(data['success'])
        self.assertIn('Required academic parameters are missing', data['error'])
        
        # Fill in the academic parameters and retry
        legacy_student_missing.admission_year = 2025
        legacy_student_missing.semester_name = 'Spring'
        legacy_student_missing.hall_attached = 'Non-Residential'
        legacy_student_missing.program = self.cse_program.name
        legacy_student_missing.cluster = self.cse_program.cluster.name
        legacy_student_missing.program_type = 'Bachelor'
        legacy_student_missing.save()
        
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertTrue(data['success'])
        self.assertIsNotNone(data['suggested_id'])
        self.assertEqual(len(data['suggested_id']), 16)
        self.assertTrue(data['suggested_id'].startswith('0802510005011'))

    def test_legacy_student_unlocked_fields(self):
        # Create a standard student with a 16-digit ID
        std_student = Student.objects.create(
            student_id='0802510005011001',
            student_name='Standard Student',
            is_legacy_student=False,
            admission_year=2025,
            semester_name='Spring',
            hall_attached='Non-Residential',
            program=self.cse_program.name,
            cluster=self.cse_program.cluster.name,
            program_type='Bachelor'
        )
        
        # Create a legacy student with a 9-digit ID
        legacy_student = Student.objects.create(
            student_id='190103024',
            student_name='Legacy Student',
            is_legacy_student=True,
            admission_year=2019,
            semester_name='Spring',
            hall_attached='Non-Residential',
            program=self.cse_program.name,
            cluster=self.cse_program.cluster.name,
            program_type='Bachelor'
        )
        
        # Test GET on edit view for standard student (fields should be disabled)
        url_std = reverse('edit_student', kwargs={'student_id': std_student.student_id})
        res_std = self.client.get(url_std)
        self.assertEqual(res_std.status_code, 200)
        form_std = res_std.context['form']
        self.assertEqual(form_std.fields['program'].widget.attrs.get('disabled'), 'disabled')
        self.assertEqual(form_std.fields['admission_year'].widget.attrs.get('disabled'), 'disabled')
        
        # Test GET on edit view for legacy student (fields should NOT be disabled)
        url_legacy = reverse('edit_student', kwargs={'student_id': legacy_student.student_id})
        res_legacy = self.client.get(url_legacy)
        self.assertEqual(res_legacy.status_code, 200)
        form_legacy = res_legacy.context['form']
        self.assertIsNone(form_legacy.fields['program'].widget.attrs.get('disabled'))
        self.assertIsNone(form_legacy.fields['admission_year'].widget.attrs.get('disabled'))

    def test_smart_campus_excel_export(self):
        # Trigger the SMART CAMPUS export URL
        response = self.client.get(reverse('export_students_smart_campus'))
        self.assertEqual(response.status_code, 200)

        # Parse the downloaded Excel content
        df = self.read_excel(response)

        # 1. Verify that 'current_batch' and 'current_semester' fields are removed
        self.assertNotIn('current_batch', df.columns)
        self.assertNotIn('current_semester', df.columns)

        # 2. Verify that 'batch' and 'semester_name' are still present
        self.assertIn('batch', df.columns)
        self.assertIn('semester_name', df.columns)

        # 3. Verify that the cancelled student (CSE901) is included in the exported students list
        student_ids = df['student_id'].astype(str).tolist()
        self.assertIn('CSE901', student_ids)

        # 4. Verify that semester_name is formatted correctly (e.g. combined with year, like "Spring 2025")
        # For CSE Student 01, it was 'Spring' and '2025' admission_year
        cse001_row = df[df['student_id'] == 'CSE001'].iloc[0]
        self.assertEqual(cse001_row['semester_name'], 'Spring 2025')

        # 5. Verify row styling for the cancelled student (CSE901) using openpyxl
        # Reading using openpyxl directly to verify colors
        from openpyxl import load_workbook
        wb = load_workbook(BytesIO(response.content))
        ws = wb['Students']
        
        # Row 1 is header, find the row index for CSE901
        cse901_row_idx = None
        for r in range(2, ws.max_row + 1):
            if ws.cell(row=r, column=2).value == 'CSE901':  # student_id column is 2nd column (SL is 1st)
                cse901_row_idx = r
                break
        
        self.assertIsNotNone(cse901_row_idx)
        # Check that the cells in that row have the yellow fill (hex FFF2CC)
        cell_fill = ws.cell(row=cse901_row_idx, column=2).fill
        self.assertIn(cell_fill.start_color.rgb, ['00FFF2CC', 'FFFFF2CC', 'FFF2CC'])

    def test_smart_campus_dept_wise_zip_export(self):
        import zipfile
        # Trigger the department-wise SMART CAMPUS export URL
        response = self.client.get(reverse('export_students_smart_campus_dept'))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/zip')

        # Load the ZIP file
        zip_file = zipfile.ZipFile(BytesIO(response.content))
        file_list = zip_file.namelist()

        # 1. Verify that files are generated for each program in the dataset (CSE, EEE, MBA)
        self.assertIn('CSE_smart_campus_export.xlsx', file_list)
        self.assertIn('EEE_smart_campus_export.xlsx', file_list)
        self.assertIn('MBA_smart_campus_export.xlsx', file_list)

        # 2. Extract and inspect one of the Excel sheets (CSE)
        cse_excel_data = zip_file.read('CSE_smart_campus_export.xlsx')
        df = pd.read_excel(BytesIO(cse_excel_data))

        # Verify columns are correct (no current_batch, current_semester)
        self.assertNotIn('current_batch', df.columns)
        self.assertNotIn('current_semester', df.columns)
        self.assertIn('batch', df.columns)
        self.assertIn('semester_name', df.columns)

        # Verify that the cancelled CSE student is included in the CSE export
        student_ids = df['student_id'].astype(str).tolist()
        self.assertIn('CSE901', student_ids)

        # Verify row styling for the cancelled student inside the zipped Excel
        from openpyxl import load_workbook
        wb = load_workbook(BytesIO(cse_excel_data))
        ws = wb['Students']
        
        cse901_row_idx = None
        for r in range(2, ws.max_row + 1):
            if ws.cell(row=r, column=2).value == 'CSE901':
                cse901_row_idx = r
                break
        
        self.assertIsNotNone(cse901_row_idx)
        cell_fill = ws.cell(row=cse901_row_idx, column=2).fill
        self.assertIn(cell_fill.start_color.rgb, ['00FFF2CC', 'FFFFF2CC', 'FFF2CC'])

    def test_smart_campus_dept_wise_modal_view(self):
        # Trigger the modal endpoint
        response = self.client.get(reverse('export_students_smart_campus_dept_modal'))
        self.assertEqual(response.status_code, 200)

        # 1. Verify that the modal lists the correct departments and counts in context
        departments = response.context['departments']
        self.assertEqual(len(departments), 3)  # CSE, EEE, MBA

        # CSE has 15 standard + 1 inactive + 1 cancelled + 1 legacy scope = 18 students
        cse_dept = next(d for d in departments if d['name'] == 'CSE')
        self.assertEqual(cse_dept['count'], 18)

        # EEE has 4 standard + 1 latest = 5 students
        eee_dept = next(d for d in departments if d['name'] == 'EEE')
        self.assertEqual(eee_dept['count'], 5)

        # MBA has 1 student
        mba_dept = next(d for d in departments if d['name'] == 'MBA')
        self.assertEqual(mba_dept['count'], 1)

        # 2. Verify download links are set up and override program filter
        self.assertIn('program=CSE', cse_dept['download_url'])
        self.assertIn('program=EEE', eee_dept['download_url'])
        self.assertIn('program=MBA', mba_dept['download_url'])

        # 3. Verify ZIP download URL is in context
        self.assertEqual(response.context['zip_download_url'], reverse('export_students_smart_campus_dept'))

    def test_studentship_preview_view(self):
        student = Student.objects.get(student_id='CSE001')
        response = self.client.get(reverse('studentship_preview', kwargs={'student_id': student.student_id}))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Studentship Certificate Generator')
        self.assertContains(response, 'BAUST/Admin-132/2015/')
        self.assertContains(response, student.student_name.upper())
        self.assertContains(response, student.father_name.upper())
        self.assertContains(response, student.mother_name.upper())

    def test_download_studentship_certificate_pdf_get(self):
        student = Student.objects.get(student_id='CSE001')
        response = self.client.get(reverse('download_studentship', kwargs={'student_id': student.student_id}))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/pdf')
        self.assertIn('inline; filename=Studentship_Certificate_CSE001.pdf', response['Content-Disposition'])

    def test_download_studentship_certificate_pdf_post(self):
        student = Student.objects.get(student_id='CSE001')
        data = {
            'ref_no': 'BAUST/Admin-132/2015/CUSTOM-999',
            'date': 'October 31, 2026',
            'heading': 'TO WHOM IT MAY CONCERN (CUSTOM)',
            'body_text': 'This is a custom body text for testing.',
            'signatory_name': 'TEST SIGNATORY',
            'signatory_title': 'Test Title',
            'signatory_contact': 'Mobile: 01999999999'
        }
        response = self.client.post(reverse('download_studentship', kwargs={'student_id': student.student_id}), data)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/pdf')
        self.assertIn('inline; filename=Studentship_Certificate_CSE001.pdf', response['Content-Disposition'])

    def test_unspecified_gender_filtering(self):
        unspecified_student = Student.objects.create(
            student_id='CSE999',
            student_name='Unspecified Gender Student',
            program=self.cse_program.name,
            admission_year=2025,
            cluster='Engineering & Technology',
            batch='25th',
            semester_name='Spring',
            program_type='Bachelor',
            admission_status='Active',
            gender=None,
            father_name='Unspecified Father',
            mother_name='Unspecified Mother',
            student_mobile='01799999999',
        )
        response = self.client.get(reverse('student_list') + '?gender=Unspecified')
        self.assertEqual(response.status_code, 200)
        page_students = response.context['page_obj'].paginator.object_list
        self.assertIn(unspecified_student, page_students)
        standard_student = Student.objects.get(student_id='CSE001')
        self.assertNotIn(standard_student, page_students)

    def test_additional_page_size_options(self):
        response = self.client.get(reverse('student_list') + '?per_page=150')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['per_page'], 150)
        self.assertEqual(response.context['page_obj'].paginator.per_page, 150)

    def test_same_program_migration_is_blocked(self):
        student = Student.objects.get(student_id='CSE001')
        
        # 1. Test utilities/execute_program_change_web blocks same program
        from students.utils import execute_program_change_web
        result = execute_program_change_web(
            student=student,
            new_program='CSE',
            new_cluster='Engineering & Technology',
            new_year=2026,
            new_semester='Spring',
            hall_name='Non-Residential'
        )
        self.assertFalse(result['success'])
        self.assertEqual(result['error'], 'Target program must be different from the current program.')
        
        # 2. Test view change_program blocks and redirects with message
        url = reverse('change_program', kwargs={'student_id': student.student_id})
        data = {
            'new_program': 'CSE',
            'new_cluster': 'Engineering & Technology',
            'new_year': '2026',
            'new_semester': 'Spring',
            'hall_name': 'Non-Residential',
            'notes': 'Test same program'
        }
        response = self.client.post(url, data)
        self.assertEqual(response.status_code, 302)
        # Check redirect is back to the change_program page
        self.assertRedirects(response, url)
        
        # Check message is set
        messages = list(response.wsgi_request._messages)
        self.assertEqual(len(messages), 1)
        self.assertEqual(str(messages[0]), "Illogical Migration: Student is already in program 'CSE'.")

    def test_program_migration_respects_id_mode_auto(self):
        from core.models import SystemSettings
        sys_settings = SystemSettings.objects.get_or_create(id=1)[0]
        sys_settings.id_mode = 'auto'
        sys_settings.save()
        
        student = Student.objects.get(student_id='CSE001')
        url = reverse('change_program', kwargs={'student_id': student.student_id})
        data = {
            'new_program': 'EEE',
            'new_cluster': 'Engineering & Technology',
            'new_year': '2026',
            'new_semester': 'Spring',
            'hall_name': 'Non-Residential',
            'notes': 'Auto migration test'
        }
        response = self.client.post(url, data)
        self.assertEqual(response.status_code, 302)
        
        # Check that a new ID was generated automatically and the student was migrated
        migrated_student = Student.objects.filter(old_student_id='CSE001').first()
        self.assertIsNotNone(migrated_student)
        self.assertTrue(migrated_student.student_id.startswith('080'))

    def test_program_migration_respects_id_mode_semi_auto(self):
        from core.models import SystemSettings
        sys_settings = SystemSettings.objects.get_or_create(id=1)[0]
        sys_settings.id_mode = 'semi_auto'
        sys_settings.save()
        
        student = Student.objects.get(student_id='CSE001')
        url = reverse('change_program', kwargs={'student_id': student.student_id})
        
        # 1. Post with invalid/missing serial
        data = {
            'new_program': 'EEE',
            'new_cluster': 'Engineering & Technology',
            'new_year': '2026',
            'new_semester': 'Spring',
            'hall_name': 'Non-Residential',
            'notes': 'Semi migration test',
            'student_id_serial': 'abc' # invalid
        }
        response = self.client.post(url, data)
        self.assertEqual(response.status_code, 302)
        messages = list(response.wsgi_request._messages)
        self.assertTrue(any("Error: In semi-auto mode, a 3-digit numeric serial must be provided." in str(m) for m in messages))
        
        # 2. Post with valid serial
        data['student_id_serial'] = '999'
        response = self.client.post(url, data)
        self.assertEqual(response.status_code, 302)
        
        migrated_student = Student.objects.filter(old_student_id='CSE001').first()
        self.assertIsNotNone(migrated_student)
        self.assertTrue(migrated_student.student_id.endswith('999'))

    def test_program_migration_respects_id_mode_manual(self):
        from core.models import SystemSettings
        sys_settings = SystemSettings.objects.get_or_create(id=1)[0]
        sys_settings.id_mode = 'manual'
        sys_settings.save()
        
        student = Student.objects.get(student_id='CSE001')
        url = reverse('change_program', kwargs={'student_id': student.student_id})
        
        # 1. Post with invalid manual ID
        data = {
            'new_program': 'EEE',
            'new_cluster': 'Engineering & Technology',
            'new_year': '2026',
            'new_semester': 'Spring',
            'hall_name': 'Non-Residential',
            'notes': 'Manual migration test',
            'manual_student_id': '123' # invalid length
        }
        response = self.client.post(url, data)
        self.assertEqual(response.status_code, 302)
        messages = list(response.wsgi_request._messages)
        self.assertTrue(any("Error: In manual mode, a valid 16-digit numeric student ID must be provided." in str(m) for m in messages))
        
        # 2. Post with valid manual ID
        data['manual_student_id'] = '0802610005021777'
        response = self.client.post(url, data)
        self.assertEqual(response.status_code, 302)
        
        migrated_student = Student.objects.filter(old_student_id='CSE001').first()
        self.assertIsNotNone(migrated_student)
        self.assertEqual(migrated_student.student_id, '0802610005021777')

    def test_directory_hall_filtering(self):
        # Setup student hall attachments
        student_non_res = Student.objects.get(student_id='CSE001')
        student_non_res.is_non_residential = True
        student_non_res.save()

        student_auah = Student.objects.get(student_id='CSE002')
        student_auah.is_non_residential = False
        student_auah.hall_attached = 'AUAH'
        student_auah.save()

        student_btbh = Student.objects.get(student_id='CSE003')
        student_btbh.is_non_residential = False
        student_btbh.hall_attached = 'BTBH'
        student_btbh.save()

        # 1. Filter by Non-Residential
        response = self.client.get(reverse('student_list'), {'hall': 'non_residential'})
        self.assertEqual(response.status_code, 200)
        students = list(response.context['page_obj'].paginator.object_list)
        self.assertIn(student_non_res, students)
        self.assertNotIn(student_auah, students)
        self.assertNotIn(student_btbh, students)

        # 2. Filter by Unspecified
        response = self.client.get(reverse('student_list'), {'hall': 'unspecified'})
        self.assertEqual(response.status_code, 200)
        students = list(response.context['page_obj'].paginator.object_list)
        self.assertNotIn(student_non_res, students)
        self.assertNotIn(student_auah, students)
        self.assertNotIn(student_btbh, students)
        # Check standard student remains (since they are residential but have no hall)
        student_default = Student.objects.get(student_id='CSE004')
        self.assertIn(student_default, students)

        # 3. Filter by specific hall AUAH
        response = self.client.get(reverse('student_list'), {'hall': 'AUAH'})
        self.assertEqual(response.status_code, 200)
        students = list(response.context['page_obj'].paginator.object_list)
        self.assertNotIn(student_non_res, students)
        self.assertIn(student_auah, students)
        self.assertNotIn(student_btbh, students)

        # 4. Filter by specific hall BTBH
        response = self.client.get(reverse('student_list'), {'hall': 'BTBH'})
        self.assertEqual(response.status_code, 200)
        students = list(response.context['page_obj'].paginator.object_list)
        self.assertNotIn(student_non_res, students)
        self.assertNotIn(student_auah, students)
        self.assertIn(student_btbh, students)

    def test_api_search_students_endpoint(self):
        # Search for students via API
        response = self.client.get(reverse('api_search_students'), {'q': 'Student'})
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertTrue(len(data['results']) > 0)

    def test_referred_by_student_saved_via_form(self):
        from students.models import Student
        from master_data.models import Batch, AdmissionYear, Hall
        
        # Ensure batch, year, and hall exist for form choices
        year_obj, _ = AdmissionYear.objects.get_or_create(year=2026, is_active=True)
        Batch.objects.get_or_create(name="26th", sort_order=26, admission_year=year_obj)
        Hall.objects.get_or_create(short_name="AUAH", code="01")
        
        # Create a referrer student
        referrer = Student.objects.create(
            student_id="CSE2026010100102",
            student_name="Referrer Student",
            program="Computer Science and Engineering",
            admission_year=2026,
            batch="26th",
            admission_status="Active"
        )
        
        # Post to create a student with referred_by_student set
        post_data = {
            'student_name': 'Candidate Student',
            'gender': 'Male',
            'dob': '2005-01-01',
            'blood_group': 'O+',
            'religion': 'Islam',
            'national_id': '1234567890',
            'student_mobile': '01712345678',
            'father_name': 'Father Name',
            'mother_name': 'Mother Name',
            'father_mobile': '01711111111',
            'mother_mobile': '01722222222',
            'program': 'Computer Science and Engineering',
            'admission_year': 2026,
            'semester_name': 'Spring',
            'batch': '26th',
            'hall_attached': 'AUAH',
            'referred_by_student': referrer.student_id,
            'admission_status': 'Active',
        }
        
        response = self.client.post(reverse('add_student'), post_data)
        self.assertEqual(response.status_code, 302)
        
        # Verify the candidate was created and linked to the referrer student
        candidate = Student.objects.get(student_name='Candidate Student')
        self.assertEqual(candidate.referred_by_student, referrer)

    def test_combined_reference_report_and_dashboard_statistics(self):
        from students.models import Student, ReferenceNode
        from students.reports import get_reference_intelligence
        
        # 1. Create a ReferenceNode
        emp_ref = ReferenceNode.objects.create(
            name_en="Test Employee Referrer",
            designation="Lecturer"
        )
        
        # 2. Create a Student Referrer
        stud_ref = Student.objects.create(
            student_id="CSE202699999",
            student_name="Test Student Referrer",
            program="Computer Science and Engineering",
            admission_year=2026,
            batch="26th",
            admission_status="Active"
        )
        
        # 3. Create students referring to each
        Student.objects.create(
            student_id="CSE202600001",
            student_name="Ref Candidate A",
            reference=emp_ref,
            batch="26th",
            admission_status="Active"
        )
        Student.objects.create(
            student_id="CSE202600002",
            student_name="Ref Candidate B",
            referred_by_student=stud_ref,
            batch="26th",
            admission_status="Active"
        )
        
        # 4. Check get_reference_intelligence combining both
        report_data = get_reference_intelligence()
        refs = report_data['references']
        
        # Verify that both are present in the references list
        names = [r['reference'] for r in refs]
        self.assertIn("Test Employee Referrer", names)
        self.assertIn("Test Student Referrer", names)
        
        # Verify columns designation and student_info
        emp_item = next(r for r in refs if r['reference'] == "Test Employee Referrer")
        self.assertEqual(emp_item['type'], 'Employee')
        self.assertEqual(emp_item['designation'], 'Lecturer')
        self.assertEqual(emp_item['student_info'], 'N/A')
        
        stud_item = next(r for r in refs if r['reference'] == "Test Student Referrer")
        self.assertEqual(stud_item['type'], 'Student')
        self.assertEqual(stud_item['designation'], 'N/A')
        self.assertIn("CSE202699999", stud_item['student_info'])
        
        # 5. Check dashboard view top_references context
        response = self.client.get(reverse('dashboard'))
        self.assertEqual(response.status_code, 200)
        top_refs = response.context['stats']['top_references']
        top_ref_names = [r['reference'] for r in top_refs]
        self.assertIn("Test Employee Referrer", top_ref_names)
        self.assertIn("Test Student Referrer", top_ref_names)

        # Verify the designations and student_info in top references
        emp_ref = next(r for r in top_refs if r['reference'] == "Test Employee Referrer")
        self.assertEqual(emp_ref['designation'], "Lecturer")
        self.assertEqual(emp_ref['type'], "Employee")

        stud_ref = next(r for r in top_refs if r['reference'] == "Test Student Referrer")
        self.assertEqual(stud_ref['designation'], "N/A")
        self.assertEqual(stud_ref['type'], "Student")
        self.assertIn("CSE202699999", stud_ref['student_info'])

    def test_api_update_board_info(self):
        """Test direct updating of SSC/HSC board verification details."""
        from django.urls import reverse
        
        student = self.inactive_student
        student.ssc_verified = True
        student.hsc_verified = True
        student.save()
        
        url = reverse('api_update_board_info')
        
        # 1. Non-privileged user should be denied
        self.client.force_login(self.scoped_user)
        response = self.client.post(url, {
            'student_id': student.student_id,
            'level': 'SSC',
            'school_college': 'New SSC School',
            'board': 'Dhaka',
            'year': '2020',
            'roll': '123456',
            'reg': '654321',
            'gpa': '5.0',
        })
        self.assertRedirects(response, reverse('user_profile'))
        
        # 2. Privileged user (superuser) should succeed for SSC
        self.client.force_login(self.superuser)
        response = self.client.post(url, {
            'student_id': student.student_id,
            'level': 'SSC',
            'school_college': 'New SSC School',
            'board': 'Dhaka',
            'year': '2020',
            'roll': '123456',
            'reg': '654321',
            'gpa': '5.00',
            'physics': '5.00',
            'chemistry': '4.50',
            'math': '4.00',
        })
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()['success'])
        
        # Verify db updates and unverified status
        student.refresh_from_db()
        self.assertEqual(student.ssc_school, 'New SSC School')
        self.assertEqual(student.ssc_board, 'Dhaka')
        self.assertEqual(student.ssc_year, '2020')
        self.assertEqual(student.ssc_roll, '123456')
        self.assertEqual(student.ssc_reg, '654321')
        self.assertEqual(student.ssc_gpa, 5.0)
        self.assertEqual(student.ssc_physics, 5.0)
        self.assertEqual(student.ssc_chemistry, 4.5)
        self.assertEqual(student.ssc_math, 4.0)
        self.assertFalse(student.ssc_verified) # Reset verified flag
        
        # 3. Privileged user should succeed for HSC
        response = self.client.post(url, {
            'student_id': student.student_id,
            'level': 'HSC',
            'school_college': 'New HSC College',
            'board': 'Rajshahi',
            'year': '2022',
            'roll': '987654',
            'reg': '123456',
            'gpa': '4.80',
        })
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()['success'])
        
        student.refresh_from_db()
        self.assertEqual(student.hsc_college, 'New HSC College')
        self.assertEqual(student.hsc_board, 'Rajshahi')
        self.assertEqual(student.hsc_year, '2022')
        self.assertEqual(student.hsc_roll, '987654')
        self.assertEqual(student.hsc_reg, '123456')
        self.assertEqual(student.hsc_gpa, 4.8)
        self.assertFalse(student.hsc_verified)



class ReferenceNodeSystemTests(TestCase):
    def setUp(self):
        from django.contrib.auth.models import User
        from core.models import Role, RolePermission
        
        self.superuser = User.objects.create_superuser(
            username='admin_ref',
            email='admin_ref@example.com',
            password='password123',
        )
        # Enable permissions for student module
        scoped_role, _ = Role.objects.get_or_create(name='Ref Staff')
        RolePermission.objects.get_or_create(role=scoped_role, module='students', task='view_directory')
        self.superuser.profile.role = scoped_role
        self.superuser.profile.save()
        self.client.force_login(self.superuser)

    def test_reference_node_auto_id_generation(self):
        from students.models import ReferenceNode
        ref1 = ReferenceNode.objects.create(name_en="First Reference")
        ref2 = ReferenceNode.objects.create(name_en="Second Reference")
        self.assertTrue(ref1.reference_id.startswith("REF-"))
        self.assertTrue(ref2.reference_id.startswith("REF-"))
        self.assertNotEqual(ref1.reference_id, ref2.reference_id)

    def test_api_search_references_endpoint(self):
        from students.models import ReferenceNode
        ref1 = ReferenceNode.objects.create(name_en="John Doe", name_bn="জন ডো", baust_id="BAUST-111")
        ReferenceNode.objects.create(name_en="Jane Smith", designation="Professor")
        
        # Search by English Name
        response = self.client.get(reverse('api_search_references'), {'q': 'John'})
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(len(data['results']), 1)
        self.assertEqual(data['results'][0]['id'], ref1.id)

        # Search by Bangla Name
        response = self.client.get(reverse('api_search_references'), {'q': 'জন'})
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(len(data['results']), 1)

        # Search by BAUST ID
        response = self.client.get(reverse('api_search_references'), {'q': '111'})
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(len(data['results']), 1)

    def test_api_reference_link_count(self):
        from students.models import ReferenceNode, Student

        # Create a reference node with no students
        ref = ReferenceNode.objects.create(name_en="Unlinked Ref", designation="Lecturer")

        # 1. Count should be 0 when no students linked
        response = self.client.get(reverse('api_reference_link_count', args=[ref.id]))
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data['linked_students'], 0)
        self.assertEqual(data['name'], "Unlinked Ref")

        # 2. Link two students and verify count
        Student.objects.create(student_id="CSE202600900", student_name="Link A", reference=ref, admission_status="Active")
        Student.objects.create(student_id="CSE202600901", student_name="Link B", reference=ref, admission_status="Active")

        response = self.client.get(reverse('api_reference_link_count', args=[ref.id]))
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data['linked_students'], 2)

        # 3. 404 for non-existent node
        response = self.client.get(reverse('api_reference_link_count', args=[999999]))
        self.assertEqual(response.status_code, 404)

    def test_api_align_legacy_reference(self):
        from students.models import ReferenceNode, Student

        # Setup standard target ReferenceNode
        target_ref = ReferenceNode.objects.create(name_en="Head of AIS", designation="Employee")

        # Setup students with legacy references
        s1 = Student.objects.create(student_id="CSE202611111", student_name="Legacy A", reference_legacy="Head AIS", admission_status="Active")
        s2 = Student.objects.create(student_id="CSE202622222", student_name="Legacy B", reference_legacy="Head AIS", admission_status="Active")
        s3 = Student.objects.create(student_id="CSE202633333", student_name="Other Ref", reference_legacy="Random Text", admission_status="Active")

        # Ensure reference field is originally empty
        self.assertNil = lambda x: self.assertIsNone(x)
        self.assertNil(s1.reference)
        self.assertNil(s2.reference)

        # Login admin user to authenticate edit access
        self.client.force_login(self.superuser)

        # Perform bulk alignment POST call
        response = self.client.post(reverse('api_align_legacy_reference'), {
            'legacy_text': 'Head AIS',
            'target_node_id': target_ref.id
        })
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertTrue(data['success'])
        self.assertEqual(data['updated_count'], 2)

        # Verify database updates
        s1.refresh_from_db()
        s2.refresh_from_db()
        s3.refresh_from_db()
        self.assertEqual(s1.reference, target_ref)
        self.assertEqual(s2.reference, target_ref)
        self.assertNil(s3.reference)

    def test_api_reference_students(self):
        from students.models import ReferenceNode, Student

        # 1. Setup a standard ReferenceNode
        ref = ReferenceNode.objects.create(name_en="Test Referrer Node", category="Employee")
        
        # 2. Setup a referring student
        referrer_student = Student.objects.create(student_id="CSE202688888", student_name="Student Referrer", admission_status="Active")

        # 3. Setup referred students
        s1 = Student.objects.create(student_id="CSE202600001", student_name="Student A", reference=ref, admission_status="Active")
        s2 = Student.objects.create(student_id="CSE202600002", student_name="Student B", referred_by_student=referrer_student, admission_status="Active")

        self.client.force_login(self.superuser)

        # 4. Query ReferenceNode referred list
        response = self.client.get(reverse('api_reference_students') + f"?type=Employee&id={ref.id}")
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Student A")
        self.assertNotContains(response, "Student B")

        # 5. Query Student referred list
        response = self.client.get(reverse('api_reference_students') + f"?type=Student&id={referrer_student.student_id}")
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Student B")
        self.assertNotContains(response, "Student A")

    def test_reference_hub_spa_actions(self):
        from students.models import ReferenceNode
        ref = ReferenceNode.objects.create(name_en="Original Name", designation="Lecturer")

        # 1. Create via HTMX
        response = self.client.post(reverse('reference_manage'), {
            'action': 'create',
            'name_en': 'New HTMX Node',
            'designation': 'Assistant Professor',
            'mobile': '01899999999'
        }, HTTP_HX_REQUEST='true')
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'New HTMX Node')
        self.assertTrue(ReferenceNode.objects.filter(name_en='New HTMX Node').exists())

        # 2. Update via HTMX
        response = self.client.post(reverse('reference_manage'), {
            'action': 'update',
            'id': ref.id,
            'name_en': 'Updated Name',
            'designation': 'Senior Lecturer'
        }, HTTP_HX_REQUEST='true')
        self.assertEqual(response.status_code, 200)
        ref.refresh_from_db()
        self.assertEqual(ref.name_en, 'Updated Name')
        self.assertEqual(ref.designation, 'Senior Lecturer')

        # 3. Delete via HTMX
        response = self.client.post(reverse('reference_manage'), {
            'action': 'delete',
            'id': ref.id
        }, HTTP_HX_REQUEST='true')
        self.assertEqual(response.status_code, 200)
        self.assertFalse(ReferenceNode.objects.filter(id=ref.id).exists())

    def test_export_excel_and_pdf_formats(self):
        from students.models import ReferenceNode
        ReferenceNode.objects.create(name_en="Excel Target Node", mobile="01799999999")

        # Excel template download
        response = self.client.get(reverse('reference_manage'), {'export': 'template'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

        # Excel export list
        response = self.client.get(reverse('reference_manage'), {'export': 'excel'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

        # PDF list
        response = self.client.get(reverse('reference_manage'), {'export': 'pdf'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/pdf')

    def test_merge_reference_nodes(self):
        from students.models import ReferenceNode, Student
        
        # Create source and target reference nodes
        source = ReferenceNode.objects.create(name_en="Ruhulamin", designation="Lecturer", mobile="01711111111")
        target = ReferenceNode.objects.create(name_en="Save Ruhul", designation="Assistant Professor", mobile="01722222222")
        
        # Create a student linked to source reference node
        student = Student.objects.create(
            student_id="CSE2026010100101",
            student_name="Test Student",
            program="Computer Science and Engineering",
            admission_year=2026,
            batch="26th",
            admission_status="Active",
            reference=source
        )
        
        # Verify initial linkage
        self.assertEqual(source.students.count(), 1)
        self.assertEqual(target.students.count(), 0)
        
        # Post request to merge source into target
        response = self.client.post(reverse('reference_manage'), {
            'action': 'merge',
            'source_id': source.id,
            'target_id': target.id
        }, HTTP_HX_REQUEST='true')
        
        self.assertEqual(response.status_code, 200)
        
        # Verify that source node is deleted and student is transferred to target
        self.assertFalse(ReferenceNode.objects.filter(id=source.id).exists())
        self.assertTrue(ReferenceNode.objects.filter(id=target.id).exists())
        
        student.refresh_from_db()
        self.assertEqual(student.reference, target)
        self.assertEqual(target.students.count(), 1)

    def test_import_references_with_smart_matching(self):
        from students.models import ReferenceNode
        import io
        import pandas as pd
        
        # Create an existing node
        existing = ReferenceNode.objects.create(
            name_en="Original Name", 
            baust_id="BAUST-1001", 
            mobile="01799999999", 
            designation="Teacher"
        )
        
        # Create an in-memory excel sheet
        df = pd.DataFrame([
            {
                'English Name': 'Original Name', 
                'BAUST ID': 'BAUST-1001',
                'Bangla Name': 'বাংলা নাম',
                'Designation': 'Updated Designation', 
                'Mobile': '01799999999'
            },
            {
                'English Name': 'Brand New Name', 
                'BAUST ID': 'BAUST-2002',
                'Bangla Name': '',
                'Designation': 'New Designation',
                'Mobile': '01788888888'
            }
        ])
        
        excel_file = io.BytesIO()
        df.to_excel(excel_file, index=False)
        excel_file.seek(0)
        excel_file.name = "import_test.xlsx"
        
        # Call the import action
        response = self.client.post(reverse('reference_manage'), {
            'action': 'import',
            'excel_file': excel_file
        })
        
        self.assertEqual(response.status_code, 302) 
        
        # Verify that "Original Name" was updated rather than duplicated
        self.assertEqual(ReferenceNode.objects.filter(name_en="Original Name").count(), 1)
        existing.refresh_from_db()
        self.assertEqual(existing.designation, "Updated Designation")
        self.assertEqual(existing.name_bn, "বাংলা নাম")
        
        # Verify that "Brand New Name" was created
        self.assertTrue(ReferenceNode.objects.filter(name_en="Brand New Name").exists())

    def test_reference_node_category_handling(self):
        from students.models import ReferenceNode
        
        # 1. Test create with default Employee category via POST
        post_data = {
            'action': 'create',
            'name_en': 'Post Employee Ref',
            'designation': 'Lecturer',
            'category': 'Employee'
        }
        response = self.client.post(reverse('reference_manage'), post_data)
        self.assertEqual(response.status_code, 200)
        node = ReferenceNode.objects.get(name_en='Post Employee Ref')
        self.assertEqual(node.category, 'Employee')
        
        # 2. Test create with External category via POST
        post_data = {
            'action': 'create',
            'name_en': 'Post External Ref',
            'designation': 'Sponsor',
            'category': 'External'
        }
        response = self.client.post(reverse('reference_manage'), post_data)
        self.assertEqual(response.status_code, 200)
        node2 = ReferenceNode.objects.get(name_en='Post External Ref')
        self.assertEqual(node2.category, 'External')
        
        # 3. Test update category via POST
        update_data = {
            'action': 'update',
            'id': node2.id,
            'name_en': 'Post External Ref',
            'category': 'Employee'
        }
        response = self.client.post(reverse('reference_manage'), update_data)
        self.assertEqual(response.status_code, 200)
        node2.refresh_from_db()
        self.assertEqual(node2.category, 'Employee')

    def test_cancellation_hub_access_and_filtering(self):
        from core.models import Role, RolePermission
        from students.models import Student
        from django.contrib.auth.models import User
        
        # Create a user with view permission
        viewer_role = Role.objects.create(name='Cancellation Viewer')
        RolePermission.objects.create(role=viewer_role, module='students', task='view_cancellations')
        
        viewer_user = User.objects.create_user(username='viewer_cancel', password='password123')
        viewer_user.profile.role = viewer_role
        viewer_user.profile.save()
        
        # Create a user without view permission
        unprivileged_user = User.objects.create_user(username='no_cancel', password='password123')
        
        # 1. Unprivileged user gets forbidden/redirected
        self.client.force_login(unprivileged_user)
        response = self.client.get(reverse('cancellation_hub'))
        self.assertEqual(response.status_code, 302) # Redirect to permission denied
        
        # 2. Privileged viewer user can access hub
        self.client.force_login(viewer_user)
        response = self.client.get(reverse('cancellation_hub'))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Cancellations Directory')
        
        # 3. Create a cancelled student and check filters
        cancelled_student = Student.objects.create(
            student_id="CSE2025999",
            student_name="Cancelled Tester",
            program="Computer Science and Engineering",
            batch="25th",
            admission_status="Cancelled"
        )
        # Create history log
        from students.models import AdmissionStatusHistory
        AdmissionStatusHistory.objects.create(
            student=cancelled_student,
            old_status="Active",
            new_status="Cancelled",
            reason_category="Migration",
            custom_notes="Going to another university",
            performed_by=self.superuser
        )
        
        # Access search/filters via HTMX
        response = self.client.get(reverse('cancellation_hub') + "?search_cancelled=Tester&target=cancelled", HTTP_HX_REQUEST='true')
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "CSE2025999")
        self.assertContains(response, "Cancelled Tester")
        
        # Filter by program
        response = self.client.get(reverse('cancellation_hub') + "?program=Computer+Science+and+Engineering&target=cancelled", HTTP_HX_REQUEST='true')
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "CSE2025999")
        
        # Filter by non-existent program
        response = self.client.get(reverse('cancellation_hub') + "?program=NonExistent&target=cancelled", HTTP_HX_REQUEST='true')
        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, "CSE2025999")
        
        # 4. Check export permission
        # Viewer doesn't have export permission
        response = self.client.get(reverse('export_cancellations_dynamic'))
        self.assertEqual(response.status_code, 302)
        
        # Give viewer export permission
        RolePermission.objects.create(role=viewer_role, module='students', task='export_cancellations')
        response = self.client.get(reverse('export_cancellations_dynamic') + "?program=Computer+Science+and+Engineering")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    def test_migration_hub_access_and_filtering(self):
        from core.models import Role, RolePermission
        from students.models import Student, ProgramChangeHistory
        from django.contrib.auth.models import User
        
        # Create a user with view permission
        viewer_role = Role.objects.create(name='Migration Viewer')
        RolePermission.objects.create(role=viewer_role, module='students', task='view_migrations')
        
        viewer_user = User.objects.create_user(username='viewer_mig', password='password123')
        viewer_user.profile.role = viewer_role
        viewer_user.profile.save()
        
        # Create a user without view permission
        unprivileged_user = User.objects.create_user(username='no_mig', password='password123')
        
        # 1. Unprivileged user gets redirected
        self.client.force_login(unprivileged_user)
        response = self.client.get(reverse('migration_center'))
        self.assertEqual(response.status_code, 302)
        
        # 2. Privileged viewer user can access hub
        self.client.force_login(viewer_user)
        response = self.client.get(reverse('migration_center'))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Migration Log History')
        
        # 3. Create a migration log history and check filters
        history_item = ProgramChangeHistory.objects.create(
            old_student_id="CSE2025001",
            new_student_id="EEE2025001",
            old_program="Computer Science and Engineering",
            new_program="Electrical and Electronic Engineering",
            notes="Interested in electrical fields"
        )
        
        # Access search/filters via HTMX
        response = self.client.get(reverse('migration_center') + "?search_history=CSE2025001&target=history", HTTP_HX_REQUEST='true')
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "CSE2025001")
        self.assertContains(response, "EEE2025001")
        
        # Filter by old program
        response = self.client.get(reverse('migration_center') + "?old_program=Computer+Science+and+Engineering&target=history", HTTP_HX_REQUEST='true')
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "CSE2025001")
        
        # Filter by non-matching old program
        response = self.client.get(reverse('migration_center') + "?old_program=Electrical+and+Electronic+Engineering&target=history", HTTP_HX_REQUEST='true')
        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, "CSE2025001")
        
        # 4. Check export permission
        # Viewer doesn't have export permission
        response = self.client.get(reverse('export_migrations_dynamic'))
        self.assertEqual(response.status_code, 302)
        
        # Give viewer export permission
        RolePermission.objects.create(role=viewer_role, module='students', task='export_migrations')
        response = self.client.get(reverse('export_migrations_dynamic') + "?old_program=Computer+Science+and+Engineering")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')










